"""Recalculate the Excel workbook with LibreOffice and compare to the Python library.

Skipped when LibreOffice (soffice) is not installed.
"""

import importlib.util
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from sdesk import cre, fixed_income as fi, hotel, tvm

ROOT = Path(__file__).resolve().parents[1]
pytestmark = pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")


@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    spec = importlib.util.spec_from_file_location("build_excel", ROOT / "scripts/build_excel.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    d = tmp_path_factory.mktemp("xl")
    src = mod.build(d / "src" / "core.xlsx")
    out = d / "out"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=180)
    return load_workbook(out / "core.xlsx", data_only=True)


R = pytest.approx


def test_loan_sheet(wb):
    ws = wb["Loan"]
    assert ws["B10"].value == R(tvm.payment(8_400_000, 0.065, 30), rel=1e-10)
    assert ws["B13"].value == R(tvm.loan_balance(8_400_000, 0.065, 30, 120), rel=1e-10)
    assert ws["B14"].value == R(ws["B13"].value, rel=1e-12)
    assert ws["B15"].value == R(tvm.effective_annual_rate(0.065, 12), rel=1e-12)
    assert ws["B378"].value == R(8_400_000)
    assert abs(ws["G376"].value) < 1e-4


def test_debt_sizing_sheet(wb):
    ws = wb["DebtSizing"]
    py = cre.max_loan(1_000_000, 16_000_000, 0.07, 30, 0.65, 1.25, 0.09)
    assert ws["B15"].value == R(py["max_loan"], rel=1e-10)
    assert ws["B16"].value.lower().startswith(py["binding"])
    assert ws["B17"].value == R(1.25)


def test_loaded_cap_sheet(wb):
    ws = wb["LoadedCap"]
    assert ws["B11"].value == R(cre.loaded_cap_value(607_927.24, 0.095, 0.023 * 0.85), rel=1e-10)
    assert ws["B14"].value == R(0.095)
    assert round(ws["B16"].value) == 268_986


def test_hotel_sheet(wb):
    ws = wb["HotelRevenue"]
    assert ws["B10"].value == R(hotel.room_revenue(120, 0.72, 125))
    assert ws["B12"].value == R(164_250)


def test_bond_sheet(wb):
    ws = wb["Bond"]
    assert ws["B9"].value == R(fi.bond_price(0.045, 0.04, 10), rel=1e-10)
    assert ws["B16"].value == R(ws["B9"].value, rel=1e-10)
    assert ws["B17"].value == R(fi.macaulay_duration(0.045, 0.04, 10), rel=1e-10)
    assert ws["B18"].value == R(fi.modified_duration(0.045, 0.04, 10), rel=1e-10)
    # built-in DURATION: agreement to 0.01 years (LibreOffice differs slightly from the formula)
    assert ws["B10"].value == R(fi.macaulay_duration(0.045, 0.04, 10), abs=0.01)
    assert ws["B14"].value == R(fi.dv01(0.045, 0.04, 10), rel=1e-6)


def test_ch01_sheet(wb):
    ws = wb["Ch01_Scenarios"]
    assert ws["B13"].value == R(2_077_700)
    assert ws["B14"].value == R(1_144_809.55, abs=0.01)
    assert round(ws["B15"].value, 4) == 0.1269
    assert round(ws["B19"].value, 4) == 0.1290
